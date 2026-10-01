"""Gera o relatório da análise em Excel, com tabelas e gráficos."""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

MOEDA = '"R$" #,##0.00'
KW = '0 "kW"'
NEGRITO = Font(bold=True)
TITULO = Font(bold=True, size=14)
CABECALHO = PatternFill("solid", fgColor="DDEBF7")
DESTAQUE = PatternFill("solid", fgColor="FCE4D6")


def _cabecalho(aba, linha, titulos):
    for col, titulo in enumerate(titulos, start=1):
        celula = aba.cell(row=linha, column=col, value=titulo)
        celula.font = NEGRITO
        celula.fill = CABECALHO
        celula.alignment = Alignment(wrap_text=True, vertical="center")


def _larguras(aba, larguras):
    for col, largura in zip("ABCDEFGH", larguras):
        aba.column_dimensions[col].width = largura


def gerar_excel(resultados, meses, parametros):
    """resultados: {nome do posto: dict de calculo.analisar_posto}. Retorna bytes do .xlsx."""
    wb = Workbook()

    # ---- Resumo
    aba = wb.active
    aba.title = "Resumo"
    aba["A1"] = "Análise de demanda contratada"
    aba["A1"].font = TITULO
    aba["A2"] = f"Gerado em {date.today():%d/%m/%Y}"
    info = [
        ("Modalidade tarifária", parametros["modalidade"]),
        ("Período analisado", f"{meses[0]} a {meses[-1]} ({len(meses)} meses)"),
        ("Crescimento de carga considerado", f"{parametros['crescimento']:.1f}%"),
    ]
    linha = 4
    for rotulo, valor in info:
        aba.cell(row=linha, column=1, value=rotulo).font = NEGRITO
        aba.cell(row=linha, column=2, value=valor)
        linha += 1

    linha += 1
    _cabecalho(aba, linha, ["Posto", "Contratada atual", "Recomendada",
                            "Custo atual no período", "Custo com a recomendada",
                            "Economia estimada", "Multas de ultrapassagem hoje"])
    for nome, r in resultados.items():
        linha += 1
        valores = [nome, r["atual"], r["otima"], r["custo_atual"], r["custo_otimo"],
                   r["economia"], r["multa_atual"]]
        formatos = [None, KW, KW, MOEDA, MOEDA, MOEDA, MOEDA]
        for col, (valor, formato) in enumerate(zip(valores, formatos), start=1):
            celula = aba.cell(row=linha, column=col, value=valor)
            if formato:
                celula.number_format = formato

    linha += 2
    aba.cell(row=linha, column=1, value="Como calculamos").font = NEGRITO
    notas = [
        "Demanda faturada = maior valor entre a contratada e a medida.",
        "Se a medida passar de 105% da contratada, o excesso é cobrado em dobro (ultrapassagem).",
        "A demanda recomendada é a que resulta no menor custo total no período analisado.",
        "Valores de custo consideram apenas a parcela de demanda da conta.",
        "Antes de pedir a alteração, confirme com a distribuidora os prazos e condições do contrato.",
    ]
    for nota in notas:
        linha += 1
        aba.cell(row=linha, column=1, value="• " + nota)
    _larguras(aba, [34, 16, 14, 20, 22, 18, 22])
    aba.row_dimensions[8].height = 32

    # ---- Uma aba de detalhe por posto
    for nome, r in resultados.items():
        det = wb.create_sheet(f"Mês a mês - {nome}"[:31])
        _cabecalho(det, 1, ["Mês", "Demanda medida (kW)",
                            f"Custo com contratada atual ({r['atual']:.0f} kW)",
                            "Multa atual",
                            f"Custo com recomendada ({r['otima']} kW)",
                            "Multa recomendada"])
        for i, (mes, a, o) in enumerate(zip(meses, r["detalhe_atual"], r["detalhe_otimo"]), start=2):
            valores = [mes, a["medida"], a["custo"], a["multa"], o["custo"], o["multa"]]
            formatos = ["@", "0.0", MOEDA, MOEDA, MOEDA, MOEDA]
            for col, (valor, formato) in enumerate(zip(valores, formatos), start=1):
                celula = det.cell(row=i, column=col, value=valor)
                celula.number_format = formato
                if a["multa"] > 0:
                    celula.fill = DESTAQUE
        total = len(meses) + 2
        det.cell(row=total, column=1, value="Total").font = NEGRITO
        for col, chave in ((3, "custo_atual"), (5, "custo_otimo")):
            celula = det.cell(row=total, column=col, value=r[chave])
            celula.number_format = MOEDA
            celula.font = NEGRITO
        _larguras(det, [10, 14, 22, 14, 22, 16])
        det.row_dimensions[1].height = 45

        barras = BarChart()
        barras.title = "Demanda medida por mês"
        barras.y_axis.title = "kW"
        barras.add_data(Reference(det, min_col=2, min_row=1, max_row=len(meses) + 1), titles_from_data=True)
        barras.set_categories(Reference(det, min_col=1, min_row=2, max_row=len(meses) + 1))
        barras.width, barras.height = 18, 8
        det.add_chart(barras, "H2")

        # Curva de custo
        det.cell(row=1, column=20, value="Demanda contratada (kW)").font = NEGRITO
        det.cell(row=1, column=21, value="Custo no período (R$)").font = NEGRITO
        for i, (d, c) in enumerate(r["curva"], start=2):
            det.cell(row=i, column=20, value=d)
            det.cell(row=i, column=21, value=c).number_format = MOEDA
        curva = LineChart()
        curva.title = "Custo total conforme a demanda contratada"
        curva.x_axis.title = "kW"
        curva.y_axis.title = "R$"
        curva.add_data(Reference(det, min_col=21, min_row=1, max_row=len(r["curva"]) + 1), titles_from_data=True)
        curva.set_categories(Reference(det, min_col=20, min_row=2, max_row=len(r["curva"]) + 1))
        curva.x_axis.tickLblSkip = 20
        curva.width, curva.height = 18, 8
        det.add_chart(curva, "H20")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
